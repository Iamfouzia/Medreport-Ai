import PyPDF2
from docx import Document
import openpyxl

def extract_text_from_pdfs(uploaded_files) -> str:
    """Extract text from PDF, Word (.docx), and Excel (.xlsx) files."""
    all_text = ""

    for file in uploaded_files:
        try:
            name = file.name.lower()

            # Step 1: Extract from PDF
            if name.endswith(".pdf"):
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        all_text += text + "\n"

            # Step 2: Extract from Word document
            elif name.endswith(".docx"):
                doc = Document(file)
                for para in doc.paragraphs:
                    if para.text.strip():
                        all_text += para.text + "\n"

            # Step 3: Extract from Excel spreadsheet
            elif name.endswith(".xlsx"):
                wb = openpyxl.load_workbook(file)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = " ".join(
                            str(cell) for cell in row if cell is not None
                        )
                        if row_text.strip():
                            all_text += row_text + "\n"

        except Exception as e:
            print(f"[ERROR] Could not read {file.name}: {e}")

    return all_text


def split_text_into_chunks(text: str, chunk_size: int = 500, overlap: int = 50) -> list:
    """Split large text into overlapping chunks for vector search."""
    chunks = []
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start = end - overlap  # Overlap keeps context between chunks

    return chunks